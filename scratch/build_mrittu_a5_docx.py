# -*- coding: utf-8 -*-
import os
import sys
import re
import shutil

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

import docx
from docx.shared import Mm, Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

FONT_BENGALI = "Kalpurush"
FONT_ARABIC = "Traditional Arabic"

COLOR_PRIMARY = RGBColor(30, 58, 138)     # Navy #1E3A8A
COLOR_SECONDARY = RGBColor(29, 78, 216)   # Blue #1D4ED8
COLOR_TEXT = RGBColor(15, 23, 42)         # Slate 900 #0F172A
COLOR_MUTED = RGBColor(100, 116, 139)     # Slate 500 #64748B
COLOR_EMERALD = RGBColor(4, 120, 87)      # Emerald 700 #047857
COLOR_WHITE = RGBColor(255, 255, 255)

def set_run_font(run, font_name=FONT_BENGALI, size_pt=10, bold=False, italic=False, color_rgb=COLOR_TEXT):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb
    
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>')
    rPr.append(rFonts)

def add_styled_paragraph(doc, text="", style='Normal', space_before=0, space_after=4, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.JUSTIFY, keep_with_next=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.alignment = align
    p.paragraph_format.keep_with_next = keep_with_next
    if text:
        add_formatted_runs(p, text)
    return p

def add_formatted_runs(p, text, base_font=FONT_BENGALI, base_size=10, base_color=COLOR_TEXT):
    # Split text by bold/italic/links/arabic
    # Simple formatting handler
    tokens = re.split(r'(\*\*[^\*]+\*\*|\*[^\*]+\*|\[[^\]]+\]\([^\)]+\)|[\u0600-\u06FF\s\(\)\.\,\:\;\!\?\"\'\«\»\۝\ۖ\ۗ\ۚ\ۛ\ۜ\۟\۠\ۡ\ۢ\ۣ\ۤ\ۥ\ۦ\ۧ\ۨ\۩\۪\۫\۬\ۭ]{6,})', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            inner = token[2:-2]
            run = p.add_run(inner)
            set_run_font(run, font_name=base_font, size_pt=base_size, bold=True, color_rgb=base_color)
        elif token.startswith('*') and token.endswith('*'):
            inner = token[1:-1]
            run = p.add_run(inner)
            set_run_font(run, font_name=base_font, size_pt=base_size, italic=True, color_rgb=base_color)
        elif token.startswith('[') and ']' in token and '(' in token and token.endswith(')'):
            m = re.match(r'\[([^\]]+)\]\(([^\)]+)\)', token)
            if m:
                label = m.group(1)
                run = p.add_run(label)
                set_run_font(run, font_name=base_font, size_pt=base_size, bold=True, color_rgb=COLOR_SECONDARY)
            else:
                run = p.add_run(token)
                set_run_font(run, font_name=base_font, size_pt=base_size, color_rgb=base_color)
        elif re.search(r'[\u0600-\u06FF]', token):
            run = p.add_run(token)
            set_run_font(run, font_name=FONT_ARABIC, size_pt=base_size+2, color_rgb=COLOR_EMERALD)
        else:
            run = p.add_run(token)
            set_run_font(run, font_name=base_font, size_pt=base_size, color_rgb=base_color)

def build_docx():
    md_path = "G:/My Drive/edu/basic-islam/মৃত্যু/মৃত্যু_জানাযা_ও_অনন্ত_জীবন_সম্পূর্ণ_বই.md"
    docx_path = "G:/My Drive/edu/basic-islam/মৃত্যু/মৃত্যু_জানাযা_ও_অনন্ত_জীবন_A5_প্রিন্ট_রেডি.docx"
    
    print(f"Reading Markdown: {md_path}")
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    doc = docx.Document()
    
    # Configure A5 Page Setup
    section = doc.sections[0]
    section.page_width = Mm(148)
    section.page_height = Mm(210)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    
    lines = content.splitlines()
    in_blockquote = False
    blockquote_lines = []
    
    for line in lines:
        stripped = line.strip()
        
        # Heading 1
        if stripped.startswith('# '):
            if blockquote_lines:
                # flush blockquote
                add_styled_paragraph(doc, " ".join(blockquote_lines), space_before=4, space_after=6)
                blockquote_lines = []
            title = stripped[2:].strip()
            p = add_styled_paragraph(doc, "", space_before=16, space_after=8, align=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True)
            run = p.add_run(title)
            set_run_font(run, font_name=FONT_BENGALI, size_pt=16, bold=True, color_rgb=COLOR_PRIMARY)
            continue
            
        # Heading 2
        if stripped.startswith('## '):
            if blockquote_lines:
                add_styled_paragraph(doc, " ".join(blockquote_lines), space_before=4, space_after=6)
                blockquote_lines = []
            h2 = stripped[3:].strip()
            p = add_styled_paragraph(doc, "", space_before=12, space_after=6, align=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True)
            run = p.add_run(h2)
            set_run_font(run, font_name=FONT_BENGALI, size_pt=13, bold=True, color_rgb=COLOR_SECONDARY)
            continue
            
        # Heading 3
        if stripped.startswith('### '):
            if blockquote_lines:
                add_styled_paragraph(doc, " ".join(blockquote_lines), space_before=4, space_after=6)
                blockquote_lines = []
            h3 = stripped[4:].strip()
            p = add_styled_paragraph(doc, "", space_before=8, space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True)
            run = p.add_run(h3)
            set_run_font(run, font_name=FONT_BENGALI, size_pt=11, bold=True, color_rgb=COLOR_TEXT)
            continue

        # Page Break
        if '<div style="page-break-before: always;"></div>' in stripped or stripped == '<div style="page-break-before: always;"></div>':
            if blockquote_lines:
                add_styled_paragraph(doc, " ".join(blockquote_lines), space_before=4, space_after=6)
                blockquote_lines = []
            doc.add_page_break()
            continue

        # Blockquote
        if stripped.startswith('>'):
            b_text = stripped.lstrip('> ').strip()
            if b_text:
                blockquote_lines.append(b_text)
            continue
        else:
            if blockquote_lines:
                p_text = " ".join(blockquote_lines)
                blockquote_lines = []
                p = add_styled_paragraph(doc, p_text, space_before=4, space_after=6)
                # add left indent for quotes
                p.paragraph_format.left_indent = Mm(4)

        if not stripped:
            continue
            
        if stripped == '---':
            continue

        # Normal paragraph
        add_styled_paragraph(doc, stripped, space_before=0, space_after=4)

    if blockquote_lines:
        add_styled_paragraph(doc, " ".join(blockquote_lines), space_before=4, space_after=6)

    print(f"Saving A5 Document to: {docx_path}")
    doc.save(docx_path)
    print("Master DOCX successfully built and saved!")

if __name__ == '__main__':
    build_docx()
